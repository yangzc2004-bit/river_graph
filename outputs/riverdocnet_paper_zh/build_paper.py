"""Create an evidence-bound Chinese manuscript with named experiments."""
from __future__ import annotations

import csv
import json
import math
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DOCX = OUT/"RiverDOCNet_密西西比河DOC重构小论文.docx"
FONTS = Path("/System/Library/Fonts/Supplemental")
EVIDENCE = {}


def records(relative):
    path = ROOT/relative
    EVIDENCE[relative] = path.stat().st_size
    with path.open() as stream:
        return list(csv.DictReader(stream))


def unique(rows, **conditions):
    found = [row for row in rows if all(row[key] == value for key, value in conditions.items())]
    if len(found) != 1:
        raise ValueError((conditions, len(found)))
    return found[0]


def fixed(value):
    return f"{float(value):.4f}"


def font(run, size=11, bold=False):
    run.font.name = "Arial Unicode MS"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:"+key), "Arial Unicode MS")
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:"+key), None)


def para(doc, text, *, style=None, size=11, bold=False, center=False):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(text), size, bold)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = Pt(size*1.3+1.2)
    if not style and not center:
        p.paragraph_format.first_line_indent = Pt(22)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    font(p.add_run(text), 14 if level == 1 else 12, True)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = Pt(20 if level == 1 else 17)
    return p


def table(doc, caption, headers, rows, widths, *, best_columns=()):
    p = para(doc, caption, size=10.5, bold=True)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    t.style = "Table Grid"
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cell is t.rows[0].cells[0] else WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(text), 10, True)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E8EEF2")
        cell._tc.get_or_add_tcPr().append(shade)
    header = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(header)
    extrema = {}
    for column in best_columns:
        numbers = [float(row[column]) for row in rows if row[column] != "—"]
        extrema[column] = max(numbers) if headers[column].startswith("R²") else min(numbers)
    for row in rows:
        cells = t.add_row().cells
        for i, (cell, text, width) in enumerate(zip(cells, row, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = Pt(14)
            bold = i in extrema and text != "—" and math.isclose(float(text), extrema[i], abs_tol=1e-10)
            font(p.add_run(text), 10, bold)
        nosplit = OxmlElement("w:cantSplit")
        t.rows[-1]._tr.get_or_add_trPr().append(nosplit)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


# Native Word mathematical objects; all notation remains editable.
def text_math(value):
    r = OxmlElement("m:r")
    token = OxmlElement("m:t")
    token.text = str(value)
    r.append(token)
    return r


def expr(*tokens):
    output = []
    for token in tokens:
        if isinstance(token, (list, tuple)):
            output.extend(token)
        else:
            output.append(text_math(token) if isinstance(token, str) else token)
    return output


def structure(tag, **parts):
    result = OxmlElement("m:"+tag)
    for name, values in parts.items():
        child = OxmlElement("m:"+name)
        for value in expr(values):
            child.append(value)
        result.append(child)
    return result


def sub(base, index):
    return structure("sSub", e=base, sub=index)


def sup(base, index):
    return structure("sSup", e=base, sup=index)


def both(base, index, power):
    return structure("sSubSup", e=base, sub=index, sup=power)


def fraction(top, bottom):
    return structure("f", num=top, den=bottom)


def root(value):
    result = OxmlElement("m:rad")
    properties = OxmlElement("m:radPr")
    hidden = OxmlElement("m:degHide")
    hidden.set(qn("m:val"), "1")
    properties.append(hidden)
    result.append(properties)
    result.append(OxmlElement("m:deg"))
    e = OxmlElement("m:e")
    for token in expr(value):
        e.append(token)
    result.append(e)
    return result


def equation(doc, *tokens):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    math_object = OxmlElement("m:oMath")
    for token in expr(*tokens):
        math_object.append(token)
    p._p.append(math_object)
    return p


def architecture():
    im = Image.new("RGB", (2200, 1060), "white")
    draw = ImageDraw.Draw(im)
    regular = ImageFont.truetype(str(FONTS/"Arial Unicode.ttf"), 37)
    small = ImageFont.truetype(str(FONTS/"Arial Unicode.ttf"), 30)
    def box(bounds, lines, color):
        draw.rounded_rectangle(bounds, radius=16, fill=color, outline="#40505c", width=3)
        x0, y0, x1, y1 = bounds
        line_height = 48
        y = (y0+y1-len(lines)*line_height)/2
        for line in lines:
            draw.text(((x0+x1)/2, y), line, fill="#14232d", font=regular, anchor="mt")
            y += line_height
    def arrow(points):
        draw.line(points, fill="#40505c", width=5, joint="curve")
        x, y = points[-1]
        ax, ay = points[-2]
        angle = math.atan2(y-ay, x-ax)
        wing = [(x, y), (x-20*math.cos(angle-.45), y-20*math.sin(angle-.45)),
                (x-20*math.cos(angle+.45), y-20*math.sin(angle+.45))]
        draw.polygon(wing, fill="#40505c")
    box((40, 165, 445, 350), ["环境基底", "15 项属性"], "#edf3e5")
    box((510, 165, 915, 350), ["环境 MLP", "24 维表示"], "#edf3e5")
    box((40, 520, 445, 715), ["水文时间序列", "12 个月 × 8 通道"], "#e7eff7")
    box((510, 520, 915, 715), ["时间 GRU", "12 维记忆"], "#e7eff7")
    box((1040, 345, 1350, 565), ["拼接与投影", "当前季节编码", "24 维节点状态"], "#f6f0e5")
    box((1450, 330, 2135, 590), ["两层方向分头图 Transformer", "上游头   下游头   全局头", "关系偏置与残差更新"], "#e5f0ef")
    box((1450, 50, 2135, 205), ["有向河网", "连接关系与方向约束"], "#f1edf6")
    box((1450, 740, 2135, 895), ["浓度读出与逆变换", "DOC 站点 × 月份矩阵"], "#f6ece5")
    arrow([(445, 258), (510, 258)])
    arrow([(445, 618), (510, 618)])
    arrow([(915, 258), (980, 258), (980, 405), (1040, 405)])
    arrow([(915, 618), (980, 618), (980, 505), (1040, 505)])
    arrow([(1350, 455), (1450, 455)])
    arrow([(1780, 205), (1780, 330)])
    arrow([(1780, 590), (1780, 740)])
    draw.text((60, 975), "DOC 观测仅用于训练监督   各模块共同训练   图中节点为具有环境与水文输入的站点", font=small, fill="#40505c")
    path = OUT/"architecture.png"
    im.save(path, dpi=(300, 300))
    return path


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    benchmark = records("experiments/phase4_transfer/doc_model_comparison_v1/summary.csv")
    relations = records("experiments/phase4_transfer/doc_river_relation_v2/analysis/summary.csv")
    ablations = records("experiments/phase4_transfer/doc_information_ablation_v1/analysis/summary.csv")
    random_stations = records("experiments/phase4_transfer/doc_unified_transfer_v2/analysis/summary.csv")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.75), Inches(.75)
    section.left_margin, section.right_margin = Inches(.8), Inches(.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial Unicode MS", Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.paragraph_format.widow_control = True
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    doc.core_properties.title = "RiverDOCNet 密西西比河流域溶解性有机碳时空重构"
    doc.core_properties.author = ""
    doc.core_properties.subject = "Environmental temporal and river-network DOC reconstruction"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    para(doc, "RiverDOCNet 密西西比河流域溶解性有机碳时空重构", style="Title", size=18, bold=True, center=True)
    heading(doc, "摘要")
    para(doc, "河流溶解性有机碳观测在时间和空间上的缺失，使不同站点的记录难以形成同步、可比较的浓度场。针对这一问题，本文提出 RiverDOCNet 河流环境时空融合网络，将环境基底编码、水文历史记忆与方向感知河网注意力整合为共同训练的重构模型。模型不以目标站点 DOC 历史作为输入，利用环境属性、过去及当前水文驱动和河网关系估计月尺度 DOC 浓度。以密西西比河流域 357 个河流站点、22,571 个观测月值为基础，设置区域站点留出、连续时间缺口和随机站点留出实验。方向分头方案在区域站点留出实验中的 MAE 为 2.1845 mg C/L，在 3、6、12 个月时间缺口中的 MAE 分别为 2.1942、2.2168、2.2866 mg C/L；随机站点留出实验的 MAE 为 1.9313 mg C/L，R² 为 0.4587。共享骨架的基础方案消融显示，环境信息提供最明确的预测贡献，时间信息亦有增益；河网贡献随评价场景变化。结果支持环境、时间与空间联合建模作为密西西比河 DOC 重构的研究框架，但尚不足以证明各组件最优、所有指标领先或任意位置连续重构的准确性。")
    para(doc, "关键词  溶解性有机碳；时空重构；环境编码；水文记忆；方向河网注意力", size=10.5)

    heading(doc, "1 引言")
    para(doc, "溶解性有机碳（dissolved organic carbon，DOC）是联系流域陆地环境与河流水体的重要碳组分。其浓度变化能够表征有机物输入、水文动员及河道内处理的综合结果，也是分析河流水质和碳输送的基础。已有研究表明，河网中的 DOC 空间分布会随水文条件改变，因此，仅用一个站点或一个采样时段的观测，难以代表整个河网的状态[1]。")
    para(doc, "DOC 重构的直接需求来自观测的不完整性。不同站点的采样日期、频率和记录长度通常不一致[2]。若上游与下游的数据来自不同月份，直接比较浓度可能将时间变化误认为空间差异；若长期记录中存在缺口，又可能遗漏季节变化和异常条件。重构的价值在于将这些离散记录组织为共同位置—时间尺度上的估计，使纵向梯度、季节差异及数据缺口能够在一致的条件下分析。本文研究浓度重构；碳通量估算还需要同步流量与额外的不确定性评估。")
    para(doc, "这一任务具有三个相互联系的难点。首先，土地覆盖、土壤及气候背景造成站点间的持久差异，相近的水文条件并不保证相同的 DOC 浓度。其次，当前浓度可能与前期水文状态有关，而水温和流量序列本身也存在缺测。最后，河网连接与地理距离并不等价，邻近站点可能位于不同支流，较远站点却可能具有真实上下游联系。仅依靠单一环境回归、局部时间插值或欧氏空间距离，未必能够同时利用这些信息。")
    para(doc, "本文据此构建 RiverDOCNet 河流环境时空融合网络，将环境基底、时间记忆和河网空间信息映射到统一节点表示，直接预测站点—月份 DOC 浓度。研究围绕三个问题展开：该框架能否在相同测试单元上取得有竞争力的重构结果；三类信息各自提供多少条件预测增益；方向分头河网编码是否优于局地预测、通用注意力和随机化河网关系。本文的证据范围是密西西比河流域月尺度重构，跨河流可复用性及任意河段位置预测留待独立研究。")

    heading(doc, "2 相关工作")
    heading(doc, "2.1 环境驱动回归与水文统计模型", 2)
    para(doc, "环境驱动模型以气候、土地覆盖、植被及水文条件建立 DOC 与流域特征的统计关系。Lu 等利用长江月尺度观测比较机器学习方法，并采用随机森林重构 DOC 变化[3]。这类方法能够处理非线性关系，是需要认真比较的强基线；但动态环境变量不等于显式历史记忆，也不等于河网消息传递。Casas-Ruiz 等利用流域面积、流量及其交互关系研究河流纵向 DOC 分布[1]，说明水文条件和沿河位置可以通过可解释的加性模型共同刻画。本文保留环境驱动随机森林、历史随机森林、历史梯度提升树和水文环境加性回归作为对标方法，不将已有多变量回归误称为单一数据源模型。")
    heading(doc, "2.2 河网空间统计与传统插值", 2)
    para(doc, "空间河网统计模型将环境回归均值与连接关系、沿河距离所决定的空间协方差结合。Wampler 等使用此类模型分析不同季节条件下的河网 DOC，并在沿河预测点上估计空间分布[4]。因此，未监测位置预测并非神经网络独有的能力。线性时间插值、形状保持插值、距离加权插值和克里金则提供计算简单或统计结构明确的参考。其实际适用性取决于可用观测的位置与日期：完全没有 DOC 记录的站点无法直接进行站内时间插值，历史稀疏月份也可能缺少同月空间邻点。本文同时记录插值直接覆盖率及补救策略，避免只在容易预测的子集上比较误差。")
    heading(doc, "2.3 过程模型与时空神经网络", 2)
    para(doc, "过程模型通过水文和碳反应过程组织 DOC 动态，例如 INCA-C 描述土壤及河道碳过程与输送[5]。这类模型的价值在于过程表达，但在新区域应用时需要相应的驱动、参数和校准条件。物理引导时空混合框架 PhyST-HC 也已结合过程状态、图神经网络和时间注意力开展有机碳预测[6]；其事件尺度任务与月尺度浓度重构的输入、目标和单位存在差异。本文的相关架构对比使用共享数据适配，而非宣称完整复现原始过程模型。")
    para(doc, "循环编码能够压缩有序历史，GRU 为紧凑的时间表示提供基础[7]。图 Transformer 则通过结构编码与注意力结合节点间关系[8]。这些方法为框架设计提供依据，但不能直接证明其在 DOC 重构中最优。本文的研究重点是稀疏 DOC 监督下三类信息的联合使用及其可检验贡献，而非将既有模块的组合宣称为首次提出的模型范式。")

    heading(doc, "3 RiverDOCNet 模型")
    heading(doc, "3.1 重构问题与输入定义", 2)
    para(doc, "将观测河网表示为有向图，节点对应具有环境与水文输入的站点，边由上游指向下游。给定站点的静态环境向量、截至目标月份的 12 个月水文窗口及河网关系，模型学习从这些输入到 DOC 浓度的映射。DOC 观测仅作为训练监督，不作为输入历史；模型也不使用站点身份嵌入或站点专属拟合参数。")
    equation(doc, sub("ŷ", "i,t"), " = ", sub("f", "θ"), "(", sub("e", "i"), ", ", sub("w", "i,t−11:t"), ", G)")
    para(doc, "环境向量包含 15 项输入：河流等级、对数汇水面积、坡度、无上游监测邻站指示量，森林、农田、城市和湿地比例，长期降水与气温、土壤有机质、海拔、基流指数，以及经纬度。无上游监测邻站指示量反映采样图结构，不等同于真实水文源头。动态窗口每月含 8 个通道：水温、对数流量、两者的有效观测标记和观测年龄，以及月份的正弦、余弦编码。观测年龄表示距离最近一次有效测量的月数，截断并缩放至 12 个月范围。归一化参数仅由训练站点拟合，动态缺测值在标准化后置零并保留标记。")
    heading(doc, "3.2 环境基底编码与时间记忆", 2)
    para(doc, "环境编码器采用两层线性映射，中间使用 GELU 激活，产生 24 维表示。它提供站点长期背景，使时间变化的解释能够依赖不同流域条件。12 个月动态窗口通过隐藏宽度为 12 的 GRU，形成目标月份的水文记忆。当前季节编码同时保留为直接输入，再与环境和时间表示拼接并投影至 24 维节点状态。环境信息不直接修改 GRU 内部更新门；三类表示也不是三个独立 DOC 预测值的平均。")
    equation(doc, sub("u", "i"), " = ", sub("W", "2"), " GELU(", sub("W", "1"), sub("e", "i"), " + ", sub("b", "1"), ") + ", sub("b", "2"))
    equation(doc, sub("g", "i,t"), " = GRU(", sub("w", "i,t−11:t"), "),   ", both("h", "i,t", "(0)"), " = ", sub("W", "f"), "[", sub("u", "i"), " ∥ ", sub("g", "i,t"), " ∥ ", sub("s", "t"), "] + ", sub("b", "f"))
    heading(doc, "3.3 方向分头河网注意力", 2)
    para(doc, "空间模块由两层图 Transformer 组成，每层设三个注意力头，每头宽度为 8。方向分头方案使一个头读取自身与上游节点，一个头读取自身与下游节点，一个头读取全部许可节点的环境水文状态。节点对关系区分自身、直接上游、其他上游祖先、直接下游、其他下游后继及无连接，并加入可学习关系偏置。已知河网中的中间站点可以用于判定路径，但不属于许可角色的中间节点，其环境、水文状态和 DOC 均不进入消息。")
    equation(doc, both("α", "ij,t", "(h)"), " = ", sub("softmax", expr("j∈", both("N", "i", "(h)"))), "(", fraction(expr(both("q", "i,t", "(h)"), "·", sup(expr("(", both("k", "j,t", "(h)"), ")"), "T")), root(sub("d", "h"))), " + ", sub("b", "h"), "(", sub("r", "ij"), "))")
    para(doc, "各头对许可节点的值向量加权求和，随后拼接投影，并通过残差连接与前馈网络更新节点表示。下游头用于统计上下文，不代表 DOC 向上游物理输送；全局头利用不同节点之间的预测关联，注意力权重也不等同于碳通量。最终读出层经过层归一化和线性映射，预测标准化的 log1p DOC，再逆变换为非负浓度；实现中对逆变换指数设定数值上限 15。完整方向分头方案具有 12,565 个可训练参数。")
    equation(doc, sub("ŷ", "i,t"), " = max{0, exp[min(15, μ + σ", sub("z", "i,t"), ")] − 1}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(architecture()), width=Inches(6.7))
    para(doc, "图 1 RiverDOCNet 方向分头方案的信息流。DOC 只用于监督；预测场由具备输入的站点和月份构成。", size=10)
    heading(doc, "3.4 训练目标与模型选择", 2)
    para(doc, "模型以训练观测上的标准化 log1p DOC 均方误差共同优化各模块。均值与标准差来自训练 DOC，月份批次的损失使用固定期望观测数作分母，使短批次不会改变单个训练观测的权重。采用 AdamW 优化器，学习率为 0.001，权重衰减为 0.0001；每批处理 32 个月，梯度范数截断为 1。最大训练轮数为 60，至少训练 15 轮，验证 MAE 连续 10 轮未改善时停止，并恢复验证最佳参数。模型和比较方案由训练及验证数据选择，不根据测试结果重新选择。")
    equation(doc, "L = ", fraction("1", expr("|", sub("Ω", "train"), "|")), sub("Σ", expr("(i,t)∈", sub("Ω", "train"))), sup(expr("[", fraction(expr("log(1+", sub("y", "i,t"), ")−μ"), "σ"), "−", sub("z", "i,t"), "]"), "2"))
    para(doc, "为使模型演进与证据一一对应，本文将早期共享骨架命名为“RiverDOCNet 联合关系注意力版”，其三个头均可读取许可节点并使用关系偏置；将当前候选命名为“RiverDOCNet 方向分头版”。基础对标和信息消融对应联合关系注意力版，方向编码和后续随机站点检验对应方向分头版。两者环境编码、GRU、表示维度及读出结构相同。方向分头版还在许可节点间保留完整已知路径；早期方案的关系来自角色诱导图，方向实验因此另设完整路径关系控制。")

    heading(doc, "4 实验设计")
    heading(doc, "4.1 研究数据与评价场景", 2)
    para(doc, "研究使用密西西比河流域河流站点队列，包含 357 个站点、324 条上游至下游的站点连接，时间范围为 1972 年 4 月至 2026 年 9 月，共 654 个月。站点来自至少 20 次 DOC 采样、记录跨度至少两年的可用队列，并限定为河流站点。DOC 观测来自 USGS，按站点—月份汇总；河网依据 NHDPlus 与 NLDI 构建，生态与流域背景属性来自 StreamCat 等数据[9]。共有 22,571 个观测月值，占 357×654 网格的 9.67%。该比例描述本队列的观测完整性，不代表全球河流的 DOC 监测覆盖率。土地覆盖与气候常年值用作静态背景，不被解释为每个历史月份的真实土地利用变化。")
    table(doc, "表 1 实验名称与检验目的", ["实验名称", "评价设置与目的"], [
        ["区域站点留出重构实验", "五个水文子区域，104 个测试站点、7,262 个观测月值；比较方法在地理缺口中的表现"],
        ["连续时间缺口重构实验", "266 个合格站点；设置 3、6、12 个月缺口，检验时间补缺能力"],
        ["组件表示选择实验", "依次比较环境编码、时间算子和融合方式，检验复杂结构是否必要"],
        ["环境时间河网消融实验", "分别移除三类信息，保持同一测试单元并重新训练，估计条件预测贡献"],
        ["河网方向编码比较实验", "比较方向分头、仅上游、路径关系与路径距离偏置方案，确定候选空间结构"],
        ["随机站点留出稳健性实验", "每次 232 个训练、54 个验证、71 个测试站点；比较真实河网、通用注意力及结构对照"],
    ], [2.0, 4.7])
    para(doc, "区域实验和随机站点实验中，测试站点的 DOC 全部不参与训练，归一化仅使用训练站点；训练图仅含训练站点，验证或测试图允许相应留出站点提供环境与水文输入，同时保留训练站点的协变量上下文。随机站点实验采用三套站点划分，每套运行三个训练种子，结果为九次指标均值。区域实验在各训练种子内汇总五个区域的测试单元后，再平均种子指标；两种划分的测试群体不同，绝对误差不能直接作升级效果比较。")
    para(doc, "时间缺口实验每个种子的评分核心固定为相同的三个月观测，扩大缺口时额外隐藏外围 DOC，但不改变评分真值。三个种子的测试单元数分别为 565、569、579。筛选要求最大缺口之外存在可用前后观测并保留足够训练数据，因此其结果适用于有时间括界条件的站点。该任务允许用缺口前后可见 DOC 训练或进行插值，是回顾性重构；神经模型的动态输入窗口仍只读取过去及当前月份的水文条件。")
    heading(doc, "4.2 对标方法与组件控制", 2)
    para(doc, "对标集合包含历史梯度提升树（XGBoost）、历史随机森林、环境驱动随机森林、水文环境加性回归（GAM）、空间河网协方差回归（SSN）和物理引导时空架构（PhyST-HC）的适配方法，以及时间和空间插值。各方法共享对应评价任务的 DOC 单位和测试单元，但按家族采用不同的输入组织和训练预算。文献适配方法表示在本数据上实现的共同任务版本，不等同于原始论文全流程复现。物理过程模型的原生示例可执行性不作为本队列上的性能结果。")
    para(doc, "组件选择实验按环境、时间、融合依次开展；后续阶段承接各区域验证所选组件，因此不同阶段的 MAE 有不同的比较父模型，不能混作同一全因子表。信息消融则以联合关系注意力版为共同父模型，在匹配输入、初始化及训练规则下重新拟合。移除环境时去掉全部静态属性和坐标；移除时间时去掉动态及季节输入；移除河网时同时取消跨站交换与显式河流描述符。另设保留静态属性的局地自路径、保留当前驱动的历史限制及无方向关系注意力控制，以区分信息包和算子作用。")
    heading(doc, "4.3 评价指标与结果解释", 2)
    para(doc, "以原始量纲 MAE 衡量总体绝对偏差，同时报告 log1p 空间 MAE，使高、低浓度区间的误差能够共同评价；RMSE 和预测 R²作为补充，反映较大误差及相对测试变异的拟合程度。MAE 和 RMSE 单位为 mg C/L，log-MAE 与 R²无量纲。R²按每个评价任务的 1−SSE/SST 计算，多个运行的 R²均值不同于汇总预测的 R²，也不同于相关系数平方。本文平均运行指标，不进行预测集成。")
    equation(doc, "MAE = ", fraction("1", "n"), sub("Σ", "k=1,…,n"), "|", sub("ŷ", "k"), "−", sub("y", "k"), "|,   ", sup("R", "2"), " = 1−", fraction(expr(sub("Σ", "k"), sup(expr("(", sub("y", "k"), "−", sub("ŷ", "k"), ")"), "2")), expr(sub("Σ", "k"), sup(expr("(", sub("y", "k"), "−ȳ)"), "2"))))
    para(doc, "所有误差在真实观测被留出的单元上计算。插值无法直接输出的单元使用事先规定的训练侧季节或总体中位数补救，报告最终完整覆盖和直接覆盖率；共同直接覆盖子集也单独检查。区域配对区间对五个区域进行聚类重采样，随机站点区间按站点成块重采样，均属于固定数据和模型下的探索性区间。重叠划分、站点水文相关及多重比较限制其推断范围。")

    heading(doc, "5 实验结果")
    heading(doc, "5.1 区域站点重构与方法对标", 2)
    names = {
        "fixed_fusion": "RiverDOCNet 联合关系注意力版",
        "doc_xgboost": "历史 XGBoost 回归",
        "ssn_covariance_adapt": "空间河网协方差 SSN 适配",
        "doc_hydro_gam": "水文环境 GAM 适配",
        "rf_history_control": "历史随机森林回归",
        "doc_ecological_rf": "环境驱动随机森林适配",
        "physt_architecture_adapt": "PhyST-HC 共享输入架构适配",
        "spatiotemporal_idw": "限窗时空距离加权插值",
        "ordinary_kriging": "指数协方差普通克里金插值",
        "euclidean_idw": "同月欧氏距离加权插值",
        "spatiotemporal_kriging": "指数协方差时空克里金插值",
        "river_idw": "同月沿河距离加权插值",
        "seasonal_median": "站点与训练侧季节中位数",
        "nearest_spatial": "同月空间最近邻插值",
        "linear_time": "站内线性时间插值",
        "pchip_time": "站内形状保持时间插值",
    }
    rows = []
    current = unique(relations, track="space", model_name="direction_heads")
    rows.append(["RiverDOCNet 方向分头版", *[fixed(current[key]) for key in ("mae", "log1p_mae", "rmse", "r2")]])
    for row in benchmark:
        if not row["space_mae"]:
            continue
        rows.append([names[row["model_name"]], *[fixed(row[key]) for key in ("space_mae", "space_log1p_mae", "space_rmse", "space_r2")]])
    table(doc, "表 2 区域站点留出重构结果", ["方法名称", "MAE ↓", "log-MAE ↓", "RMSE ↓", "R² ↑"], rows, [3.2, .85, .95, .85, .85], best_columns=(1, 2, 3, 4))
    para(doc, "联合关系注意力版在该比较中的 MAE 和 log-MAE 最低，分别为 2.1839 和 0.3085；方向分头版的对应结果为 2.1845 和 0.3112，两者空间绝对误差接近，分头版的 log-MAE 略高。方向分头版的 RMSE 为 4.6436、R²为 0.5090，而水文环境加性回归的 RMSE 为 4.5142、R²为 0.5361，说明 RiverDOCNet 不占据所有指标的最优位置。")
    para(doc, "与最低 MAE 的传统空间方法——限窗时空距离加权插值——相比，联合关系注意力版的 MAE 点估计降低 27.40%；相对历史梯度提升树和空间河网协方差回归分别降低 3.05% 和 6.13%。后两项的探索性区域配对区间包含零，故只能表述为本开发数据上的有利点估计，不能认定稳定总体优势。传统插值直接覆盖率为 99.60%～99.92%，本文表格包含规定补救；共同直接覆盖子集保留了融合模型与传统空间方法之间的排序。")

    heading(doc, "5.2 时间缺口重构与方向编码", 2)
    time_rows = []
    for arm, label in (("direction_heads", "RiverDOCNet 方向分头版"), ("fixed_fusion", "RiverDOCNet 联合关系注意力版"), ("generic_dense", "无方向约束全局注意力方案"), ("upstream_only", "仅上游注意力方案"), ("path_bias", "方向路径距离偏置方案")):
        time_rows.append([label, *[fixed(unique(relations, track="time", gap_months=str(gap), model_name=arm)["mae"]) for gap in (3, 6, 12)]])
    for name in ("linear_time", "pchip_time", "seasonal_median", "spatiotemporal_idw", "spatiotemporal_kriging"):
        row = unique(benchmark, model_name=name)
        time_rows.append([names[name], *[fixed(row[f"time_mae_{gap}"]) for gap in (3, 6, 12)]])
    table(doc, "表 3 连续时间缺口中的 MAE", ["方法名称", "3 个月 ↓", "6 个月 ↓", "12 个月 ↓"], time_rows, [3.7, 1., 1., 1.], best_columns=(1, 2, 3))
    para(doc, "方向分头版在三个缺口长度上的 MAE 分别为 2.1942、2.2168 和 2.2866，相对联合关系注意力版降低 3.79%、3.96% 和 1.98%；原始及对数绝对误差在九个缺口—种子任务中均降低。该方案依据空间与时间验证指标的预定联合规则选出，没有根据测试表中的赢家改变模型。")
    para(doc, "方向分头版并非在所有空间或时间指标上均最优。区域空间实验中，方向路径距离偏置方案具有稍低的 MAE 点估计 2.1830，而其时间缺口表现未稳定改善；这里的路径距离指站点图跳数，不是实际河道长度或输送时间。12 个月缺口中，无方向约束全局注意力的 MAE 已达到 2.2876，与分头版 2.2866 十分接近。因此，长缺口提升的一部分与广泛节点上下文有关，不能全部归因于真实上下游方向。联合关系注意力版与传统时间插值的优势也依赖指标；例如三个月缺口的线性插值在部分平方误差指标上更好。")

    heading(doc, "5.3 信息消融与组件选择", 2)
    ablation_names = {"fixed_fusion": "完整环境时间河网融合", "no_environment": "去除静态环境信息", "no_time": "去除动态与季节信息", "no_river": "去除跨站交换与河流描述符"}
    ablation_rows = []
    full = unique(ablations, track="space", model_name="fixed_fusion")
    for name, label in ablation_names.items():
        row = unique(ablations, track="space", model_name=name)
        increase = 100*(float(row["mae"])/float(full["mae"])-1)
        ablation_rows.append([label, fixed(row["mae"]), fixed(row["log1p_mae"]), fixed(row["rmse"]), fixed(row["r2"]), f"{increase:+.2f}%"])
    table(doc, "表 4 联合关系注意力版的空间信息消融", ["方案名称", "MAE ↓", "log-MAE ↓", "RMSE ↓", "R² ↑", "MAE 变化"], ablation_rows, [2.65, .75, .95, .75, .75, .85], best_columns=(1, 2, 3, 4))
    para(doc, "去除静态环境包使 MAE 增加 67.68%，去除动态及季节输入增加 4.91%，去除跨站交换与河流描述符增加 1.36%。前两项在五区域的两项绝对误差配对区间中均支持完整模型，河网移除的空间区间则包含零。保留静态属性但仅使用自路径时，MAE 增加 2.20%；仅去掉关系类别而保留通用全局注意力时，MAE 只增加 0.04%。这些对照说明静态背景的条件预测贡献最明确，不能将三项误差变化相加为 DOC 成因比例，也不能将整包环境移除解释为某个单一生态变量的因果作用。")
    para(doc, "独立的组件选择进一步限制了“最优组件”主张。环境阶段中，线性映射 MAE 为 2.1749，紧凑非线性编码为 2.1848，残差环境编码为 2.2332；线性映射与非线性编码的差异不稳定。时间阶段中，历史展开编码为 2.2028、GRU 为 2.2083、长短期记忆编码为 2.2482、因果时间注意力为 2.2254，均未建立稳定的 GRU 优势。融合阶段的拼接投影、残差拼接、环境条件融合与自适应权重融合分别为 2.1943、2.2034、2.1984、2.1909；最低点估计相对拼接投影的优势约 0.16%，区间包含零。不同阶段有不同父模型，不应将上述数字与表 2 直接排名。保留紧凑环境 MLP、GRU 和拼接投影，依据的是信息组织和实现简洁，以及当前比较中缺乏可靠替代优势，并非全局最优性证明。")
    para(doc, "空间算子选择在相同环境与时间骨架下比较了局地自路径、有向近邻图卷积、河流等级分层图卷积和联合关系图注意力。其 MAE 分别为 2.2318、2.3261、2.2961 和 2.1839，图注意力相对两种图卷积在五个区域均具有较低 MAE。注意力能够直接比较较远节点的状态，为稀疏站点图提供较广上下文；但这一比较同时改变读取范围与算子结构，不能把全部优势归于注意力机制，也不证明河流等级分层本身有稳定增益。它支持在已测试实现中优先采用图 Transformer，并通过后续方向编码与空间控制继续检验河网贡献。")

    heading(doc, "5.4 随机站点留出中的空间贡献", 2)
    random_names = {"direction_heads": "真实方向分头河网方案", "generic_dense": "无方向约束全局注意力方案", "self_only": "仅局地自路径方案", "permuted_river": "节点关系随机化河网方案"}
    random_rows = []
    for arm, label in random_names.items():
        row = unique(random_stations, domain="source_hidden_station", arm=arm)
        random_rows.append([label, *[fixed(row[key]) for key in ("mae", "log1p_mae", "rmse", "r2")]])
    table(doc, "表 5 随机站点留出稳健性结果", ["空间方案", "MAE ↓", "log-MAE ↓", "RMSE ↓", "R² ↑"], random_rows, [3.2, .85, .95, .85, .85], best_columns=(1, 2, 3, 4))
    para(doc, "在另一套三次随机站点划分中，真实方向分头方案的 MAE 为 1.9313，低于局地自路径的 2.0819 和关系随机化的 2.0593，分别降低 7.24% 和约 6.22%；两项绝对误差的探索性站点配对区间均支持真实方案。其相对通用全局注意力的 MAE 优势约为 1.70%，配对区间包含零。这支持该场景中跨站信息及真实河网关系的预测价值，但仍不足以证明分头编码优于所有通用注意力实现。关系随机化保留类别计数与结构分布，部分真实连接可能残留，因而它不是严格完全无连接的负对照。")
    para(doc, "该表与表 2 使用不同留出群体和汇总方式，不能将 2.1845 到 1.9313 解读为模型升级带来的误差下降。结果差异本身表明，空间模块贡献会随地理缺口形态和可用上下文变化；评价场景应当与最终应用中的缺测方式一致。")

    heading(doc, "6 讨论")
    para(doc, "本文的主要发现是三类信息能够在同一表示中联合学习，而其贡献强度并不均衡。环境基底承担最明确的地理背景解释，动态与季节信息使预测不再只是每站恒定浓度，河网模块则在局地表示之外提供条件上下文。消融支持信息价值，却不自动证明指定编码器的最优性。以简单拼接组织表示，有助于将更多复杂结构的收益与其代价分开检验。")
    para(doc, "方向分头设计为稀疏站点图提供了可读的关系划分，其较明确的收益目前来自时间缺口实验与随机站点控制。区域空间实验中，联合关系注意力和通用注意力仍具有竞争力，故本文将方向分头作为受验证支持的候选方案，而不将注意力权重解释为碳输送机制。整体 MAE 的有利结果也不能覆盖加性回归在 RMSE 和 R²上的优势。浓度极值与普通浓度具有不同评价权重，最终应用应共同检查原始及对数误差。")
    para(doc, "动态输入可用性是重构应用的重要边界。组件研究中，DOC 已观测测试月份的水温可用率为 95.6%，而测试站点记录跨度内自然缺测 DOC 单元仅为 32.6%；真实缺测单元也更容易缺少完整水文历史。隐藏已有 DOC 的测试，可能比实际缺测位置具备更丰富的辅助信息。因此，完整输出覆盖率不是完整准确性保证，月尺度站点预测也不能直接推广为连续时间或任意河段位置的高精度场。")
    para(doc, "本文限制在一个已多次分析的开发流域，区域重采样只有五个簇，重复种子也不增加独立流域数量。将框架用于其他河流时，可以在对应河流数据上重新训练，但其表现需要各自验证。虚拟站点应用还需提供河段环境与动态驱动、建立一致的河网查询表示，并在未参与训练的位置检验；这些是与当前月尺度站点重构相衔接的后续任务。")

    heading(doc, "7 结论")
    para(doc, "本文提出 RiverDOCNet 河流环境时空融合网络，以环境 MLP、12 个月 GRU 和方向分头图 Transformer 联合估计密西西比河流域月尺度 DOC。匹配比较表明，该框架在总体绝对误差上具有竞争力，环境信息的贡献最明确，方向分头方案改善了已有时间缺口任务，真实河网信息在随机站点留出中提供了预测增益。研究同时保留不利和不确定结果：组件最优性尚未确立，部分统计模型在平方误差指标上更好，分头结构相对通用注意力的优势依赖场景。现有证据支持该框架用于已声明输入条件下的流域内重构研究，不构成全球最优、跨河流零样本成功或任意位置连续重构的证明。")

    heading(doc, "参考文献")
    references = [
        "[1] Casas-Ruiz J P, Spencer R G M, Guillemette F, et al. Delineating the Continuum of Dissolved Organic Matter in Temperate River Networks. Global Biogeochemical Cycles, 2020, 34: e2019GB006495. https://doi.org/10.1029/2019GB006495",
        "[2] Sterle G, Perdrial J, Kincaid D W, et al. CAMELS-Chem: augmenting CAMELS with atmospheric and stream water chemistry data. Hydrology and Earth System Sciences, 2024, 28: 611–630. https://doi.org/10.5194/hess-28-611-2024",
        "[3] Lu X, Wu Y, Fan J, Ye Y, Wang J. Long-term variations and driving factors of dissolved organic carbon transport in the Yangtze River using machine learning. Journal of East China Normal University Natural Science, 2026, 3: 17–29, 43. https://doi.org/10.3969/j.issn.1000-5641.2026.03.002",
        "[4] Wampler K A, Bladon K D, Myers-Pigg A N. The influence of burn severity on dissolved organic carbon concentrations across a stream network differs based on seasonal wetness conditions. Biogeosciences, 2024, 21: 3093–3120. https://doi.org/10.5194/bg-21-3093-2024",
        "[5] Futter M N, et al. Modeling the mechanisms that control in-stream dissolved organic carbon dynamics in upland and forested catchments. Water Resources Research, 2007. https://doi.org/10.1029/2006WR004960",
        "[6] Budamala V, Choi E, Kim M, Lee S, Kim H. Hierarchical prediction of streamflow and organic carbon at event scale using a physics-guided AI hybrid framework. Journal of Hydrology, 2026: 136248. https://doi.org/10.1016/j.jhydrol.2026.136248",
        "[7] Cho K, van Merriënboer B, Gulcehre C, et al. Learning Phrase Representations using RNN Encoder–Decoder for Statistical Machine Translation. EMNLP, 2014: 1724–1734. https://doi.org/10.3115/v1/D14-1179",
        "[8] Ying C, et al. Do Transformers Really Perform Bad for Graph Representation? Advances in Neural Information Processing Systems, 2021. https://arxiv.org/abs/2106.05234",
        "[9] Hill R A, Weber M H, Leibowitz S G, Olsen A R, Thornbrugh D J. The Stream-Catchment Dataset: A Database of Watershed Metrics for the Conterminous United States. Journal of the American Water Resources Association, 2016, 52: 120–128. https://doi.org/10.1111/1752-1688.12372",
    ]
    for item in references:
        p = para(doc, item, size=10)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_together = True
    for border in list(doc.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.save(DOCX)
    with ZipFile(DOCX) as archive:
        from lxml import etree
        xml = etree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
        all_text = "\n".join(xml.xpath("//w:t/text()", namespaces=ns))
        prohibited = ("split142", "seed42", "fixed_fusion", "direction_heads", "G0", "H2X", "hash", "哈希", "permuted_river", "phase4_transfer")
        if any(token in all_text for token in prohibited):
            raise ValueError("internal run identifier leaked into paper")
        if len(xml.xpath("//m:oMath", namespaces=ns)) != 7:
            raise ValueError("expected seven native equations")
    writeout = {"deliverable": DOCX.name, "scope": "Mississippi only; directional main model with explicitly named joint-attention reference",
        "source_tables": EVIDENCE, "native_equations": 7, "characters": len(all_text),
        "tables": len(doc.tables), "references": len(references), "internal_identifiers_in_manuscript": False}
    (OUT/"evidence.json").write_text(json.dumps(writeout, indent=2, ensure_ascii=False)+"\n")
    print(json.dumps(writeout, ensure_ascii=False))


if __name__ == "__main__":
    build()
